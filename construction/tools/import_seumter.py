"""세움터 시군구별 집계 → trend.html의 시·군·구 데이터(SUB)에 반영.

세움터 「인허가정보집계」에서 받은 공식 시·군·구 수치로, 자체 집계(update_trend.py)한 값을 덮어쓴다.
덮어쓴 달은 SUB_SRC에 기록되어 화면 안내문에 '세움터 집계'로 표시된다.

입력 두 가지 (여러 개 섞어서 줄 수 있음)
  1) 세움터 월별 파일  : 「시군구별_용도별_건축허가_착공_준공_현황(YYYY년MM월).xlsx」 → 그 달 허가·착공
  2) 착공 백데이터 파일 : 시트 '데이터'(연월·구분·시도_현재·시군구_현재·…·연면적_㎡) → 여러 달 착공

사용법 (openpyxl 필요: pip install openpyxl)
  python import_seumter.py <백데이터.xlsx> <월별파일1.xlsx> [<월별파일2.xlsx> ...]
  python import_seumter.py ... --dry-run

순서: 매달 update_trend.py로 기간(DATA)을 새 달까지 민 뒤, 이 스크립트에 백데이터와 그 뒤 월별 파일을 모두 넘긴다
(update_trend.py는 SUB를 원자료로 새로 만들기 때문에 세움터 값을 다시 덮어써야 함).
"""
import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

TREND = Path(__file__).resolve().parent.parent / 'trend.html'
SIDO = {'부산광역시': '부산', '울산광역시': '울산', '대구광역시': '대구', '경상남도': '경남', '경상북도': '경북'}
DIRECT = '시·도 직접'   # 시군구 칸이 빈 행 = 시·도가 직접 처리한 건(대단지 주택 등)


def district(name):
    if not name or name in ('시도 직접', '소계'):
        return DIRECT if name != '소계' else None
    name = str(name).replace(' ', '')
    m = re.match(r'(창원시|포항시)', name)          # 일반구는 시 단위로 합침(사이트 기준)
    return m.group(1) if m else name


def read_monthly(path):
    """세움터 월별 파일 → {('허가'|'착공', 시도, 구군, 'YYYY-MM'): ㎡}"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    text = ' '.join(str(c) for r in wb['목차'].iter_rows(values_only=True) for c in r if c)
    m = re.search(r'(\d{4})\.(\d{2})\.01\s*~', text) or re.search(r'(\d{4})년\s*(\d{2})월', Path(path).name)
    if not m:
        sys.exit(f'{path}: 기준 월을 찾지 못했습니다.')
    month = f'{m.group(1)}-{m.group(2)}'
    out = defaultdict(float)
    totals = {}
    for ws in wb.worksheets:
        kind = '허가' if '허가' in ws.title else '착공' if '착공' in ws.title else None
        if not kind or not ws.title.endswith('_01'):       # _01 시트의 3번째 열이 총계
            continue
        for r in ws.iter_rows(values_only=True):
            if r and r[0] in SIDO and isinstance(r[2], (int, float)):
                if r[1] == '소계':
                    totals[(kind, SIDO[r[0]])] = r[2]
                else:
                    out[(kind, SIDO[r[0]], district(r[1]), month)] += r[2]
    for (kind, sd), v in totals.items():                  # 소계와 구·군 합이 맞는지 확인
        s = sum(x for (k, a, _, mm), x in out.items() if k == kind and a == sd)
        if abs(s - v) > 1:
            print(f'  경고 {path.name} {month} {kind} {sd}: 소계 {v:,.0f} ≠ 구·군 합 {s:,.0f}')
    print(f'  {Path(path).name}: {month} 허가·착공')
    return out


def read_backdata(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = defaultdict(float)
    for r in wb['데이터'].iter_rows(min_row=2, values_only=True):
        if r[3] == '착공' and r[4] in SIDO and r[2]:
            out[('착공', SIDO[r[4]], district(r[5]), r[2])] += r[15] or 0
    months = sorted({k[3] for k in out})
    print(f'  {Path(path).name}: 착공 {months[0]} ~ {months[-1]}')
    return out


def ranges(months):
    """['2021-09','2021-10',…] → [['2021-09','2026-08'], …] (연속 구간)"""
    def idx(m): y, mm = map(int, m.split('-')); return y * 12 + mm
    out = []
    for m in sorted(months):
        if out and idx(m) == idx(out[-1][1]) + 1:
            out[-1][1] = m
        else:
            out.append([m, m])
    return out


def corr(a, b):
    n = len(a); ma, mb = sum(a) / n, sum(b) / n
    c = sum((x - ma) * (y - mb) for x, y in zip(a, b))
    va = sum((x - ma) ** 2 for x in a) ** .5; vb = sum((y - mb) ** 2 for y in b) ** .5
    return c / (va * vb) if va and vb else 0.0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('files', nargs='+')
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args()

    html = TREND.read_text(encoding='utf-8')
    data = json.loads(re.search(r'const DATA =(\{.*?\});\n', html).group(1))
    sub = json.loads(re.search(r'const SUB = (\{.*?\});\n', html).group(1))
    months = data['months']
    n = len(months)

    # 파일 순서대로 읽고, 같은 종류·같은 달이 겹치면 뒤에 준 파일 값으로 통째로 바꿈
    vals = {}
    for f in a.files:
        p = Path(f)
        part = read_backdata(p) if '데이터' in openpyxl.load_workbook(p, read_only=True).sheetnames else read_monthly(p)
        months_in = {(k[0], k[3]) for k in part}
        for key in [k for k in vals if (k[0], k[3]) in months_in]:
            del vals[key]
        vals.update(part)

    covered = defaultdict(set)                        # 종류별로 세움터 값으로 바꾼 달
    for kind, sd, dist, m in vals:
        if m in months:
            covered[kind].add(m)

    col = {'허가': 1, '착공': 2}
    for sd in data['regions']:
        rows = {r[0]: r for r in sub[sd]}
        names = {d for (k, s, d, m) in vals if s == sd and m in months}
        for d in sorted(names - set(rows)):           # 시·도 직접 등 새 항목
            rows[d] = [d, [0] * n, [0] * n]
            sub[sd].append(rows[d])
        for kind, ms in covered.items():
            for m in ms:
                i = months.index(m)
                for r in sub[sd]:
                    r[col[kind]][i] = round(vals.get((kind, sd, r[0], m), 0))
        # 시·도 직접은 맨 뒤로
        sub[sd].sort(key=lambda r: r[0] == DIRECT)

    src = {k: ranges(v) for k, v in covered.items()}
    valid = {}
    for sd in data['regions']:
        tp = [sum(r[1][i] for r in sub[sd]) for i in range(n)]
        ts = [sum(r[2][i] for r in sub[sd]) for i in range(n)]
        op, os_ = data['series'][sd]['permit_area'], data['series'][sd]['start_area']
        valid[sd] = [round(sum(tp) / sum(op), 2), round(corr(tp, op), 2), round(sum(ts) / sum(os_), 2), round(corr(ts, os_), 2)]
    print('세움터 반영 구간:', json.dumps(src, ensure_ascii=False))
    print('시도 합계 대비 [허가 비율, 상관, 착공 비율, 상관]:', json.dumps(valid, ensure_ascii=False))
    if a.dry_run:
        return
    dump = lambda o: json.dumps(o, ensure_ascii=False, separators=(',', ':'))
    html = re.sub(r'const SUB = \{.*?\};\n', lambda _: f'const SUB = {dump(sub)};\n', html, count=1)
    html = re.sub(r'const SUB_VALID = \{.*?\};\n', lambda _: f'const SUB_VALID = {dump(valid)};\n', html, count=1)
    if 'const SUB_SRC' in html:
        html = re.sub(r'const SUB_SRC = \{.*?\};\n', lambda _: f'const SUB_SRC = {dump(src)};\n', html, count=1)
    else:
        html = html.replace('const SUB_VALID = ', f'// 세움터 시군구 집계로 바꾼 달(종류별 연속 구간). 나머지 달은 건축HUB 원자료 자체 집계\nconst SUB_SRC = {dump(src)};\nconst SUB_VALID = ', 1)
    TREND.write_text(html, encoding='utf-8', newline='')
    print(f'{TREND} 갱신 완료')


if __name__ == '__main__':
    main()
